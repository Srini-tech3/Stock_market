"""Persist dashboard-only Open/Closed selections without editing the Excel journal."""
import hashlib
import json
import math
import threading
from datetime import date, datetime
from numbers import Real
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook

from app_support import atomic_json


INITIAL_CAPITAL = 300000


def journal_kpis(headers, rows, statuses, workbook_available=True):
    """Summarize manual amounts; never evaluate formulas or use estimated risk.

    P&L totals cover all rows. Only Open rows reserve margin. Available capital
    follows the user's formula (capital + profit - used); loss is shown separately.
    Blank cells mean unrecorded/zero. Invalid values make a total unavailable.
    """
    columns = [str(header).strip().casefold() for header in headers]

    def total(column, selected, loss=False):
        if not workbook_available:
            return None
        if not selected:
            return 0.0
        if column not in columns:
            return None
        index = columns.index(column)
        amounts = []
        for row in selected:
            value = row[index] if index < len(row) else None
            if value is None or isinstance(value, str) and not value.strip():
                continue
            if isinstance(value, bool) or not isinstance(value, Real) or not math.isfinite(value):
                return None
            if not loss and value < 0:
                return None
            amounts.append(abs(float(value)) if loss else float(value))
        try:
            return round(math.fsum(amounts), 2)
        except OverflowError:
            return None

    profit = total("profit", rows)
    loss = total("loss", rows, loss=True)
    used = (None if any(status not in ("open", "closed") for status in statuses)
            else total("standalone margin", [row for row, status in zip(rows, statuses) if status == "open"]))
    available = None if profit is None or used is None else round(INITIAL_CAPITAL + profit - used, 2)
    return {"capital": INITIAL_CAPITAL, "total_profit": profit, "total_loss": loss,
            "available_capital": available, "used_capital": used}


class JournalStatusError(ValueError):
    pass


def read_journal(path):
    path = Path(path)
    if not path.exists():
        return [], []
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = workbook["Journal"]
            headers = [cell.value for cell in sheet[1]]
            rows = [row for row in sheet.iter_rows(min_row=2, values_only=True)
                    if any(value is not None for value in row)]
            return headers, rows
        finally:
            workbook.close()
    except (OSError, KeyError, ValueError, BadZipFile):
        raise ValueError("Unable to read the journal workbook. Check that the Journal sheet exists and the file is accessible.") from None


def journal_entry_ids(headers, rows):
    # Exclude P&L, manual fields and rank: changing them must not reset a selection.
    identity_columns = ("expirydate", "strategy", "option", "sellstrike", "buystrike",
                        "sellpremium", "buypremium", "estimatedpop")
    columns = [str(header).strip().casefold() for header in headers]
    names = [name for name in identity_columns if name in columns]
    if not names:
        names = ["rank"] if "rank" in columns else []

    def normalize(value):
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, Real):
            return round(float(value), 4) if math.isfinite(value) else None
        return str(value).strip() if value is not None else None

    counts, ids = {}, []
    for row in rows:
        fields = dict(zip(columns, row))
        identity = [(name, normalize(fields.get(name))) for name in names]
        digest = hashlib.sha256(json.dumps(identity, sort_keys=True).encode("utf-8")).hexdigest()
        occurrence = counts.get(digest, 0)
        counts[digest] = occurrence + 1
        ids.append(f"{digest}-{occurrence}")
    return ids


class JournalStatusStore:
    def __init__(self, output_dir):
        self.path = Path(output_dir) / "journal_status.json"
        self.lock = threading.Lock()

    def load(self):
        if not self.path.exists():
            return {}
        try:
            payload = json.loads(self.path.read_text(encoding="utf-8"))
            if not isinstance(payload, dict) or payload.get("schema_version") != 1:
                raise ValueError("Invalid schema")
            statuses = payload.get("statuses")
            if not isinstance(statuses, dict) or not all(isinstance(key, str) and value in ("open", "closed")
                                                         for key, value in statuses.items()):
                raise ValueError("Invalid statuses")
            return statuses
        except (OSError, ValueError, TypeError):
            raise JournalStatusError("Saved trade selections could not be read from Output/journal_status.json. Controls are disabled to avoid overwriting them.") from None

    def set_status(self, entry_id, status):
        if status not in ("open", "closed"):
            raise ValueError("Trade status must be open or closed.")
        with self.lock:
            statuses = self.load()
            statuses[entry_id] = status
            atomic_json(self.path, {"schema_version": 1, "statuses": statuses})
            return statuses
