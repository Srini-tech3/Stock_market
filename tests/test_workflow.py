"""Offline regression tests. Fixtures are historical samples, never live quotes."""
import hashlib
import json
import logging
import math
import re
import shutil
import sys
import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import Mock, patch

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "Script"))

from app import create_app
from config import ConfigManager
from data_provider import MarketDataError
from main import AnalysisBusy, run_analysis
from journal_state import journal_kpis
from market_context import MarketContext
from option_chain import OptionChain
from scorer import OptionScorer
from strategies import StrategyEngine
from trade_tracker import TradeTracker


RULES = {
    "SpotPrice": 23779.15, "VIX": 11.16, "Trend": "Sideways", "CPRWidth": "Wide",
    "LotSize": 65, "WingWidth": 300, "AvailableCapital": 300000, "MinPOP": 70,
    "MinDelta": .1, "MaxDelta": .4, "MinOI": 10000, "MinVolume": 5000,
    "MaxBidAskPct": 3, "MaxSpreadPoints": 2, "MaxCapitalPerTrade": 250000,
}


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.input = self.root / "Input"
        self.output = self.root / "Output"
        self.input.mkdir()
        self.output.mkdir()
        # Relative dates make tests repeatable; market values remain historical fixture data.
        self.expiry = (date.today() + timedelta(days=13)).isoformat()
        stamp = datetime.now().strftime("%Y-%m-%d-%H-%M-%S")
        self.csv = self.input / f"NIFTY_{self.expiry}_option_chain_{stamp}.csv"
        shutil.copy2(ROOT / "tests/fixtures/nifty_snapshot.csv", self.csv)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "CONFIG"
        sheet.append(["Parameter", "Value"])
        for key, value in RULES.items():
            sheet.append([key, value])
        workbook.save(self.input / "Config.xlsx")
        workbook.close()
        self.manager = ConfigManager(self.input, self.output)
        self.app = create_app(self.manager)
        self.client = self.app.test_client()

    def tearDown(self):
        for handler in logging.getLogger("option_analyzer").handlers[:]:
            logging.getLogger("option_analyzer").removeHandler(handler)
            handler.close()
        self.temporary.cleanup()

    def scan(self, source="csv"):
        return self.client.post("/run-analysis", json={"expiry": self.expiry, "market_source": source})

    def load_result(self):
        return json.loads((self.output / "result.json").read_text(encoding="utf-8"))

    def journal_ids(self):
        html = self.client.get("/journal").get_data(as_text=True)
        return re.findall(r'data-entry-id="([^"]+)"', html)

    def set_journal_status(self, entry_id, status):
        return self.client.post("/api/journal/status", json={"entry_id": entry_id, "status": status})

    def test_complete_csv_api_dashboard_workflow_without_broker(self):
        original_config = (self.input / "Config.xlsx").read_bytes()
        with patch("data_provider.fetch_spot_and_vix", side_effect=AssertionError("No broker in CSV mode")):
            response = self.scan()
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["journal_rows_added"], 4)
        data = self.load_result()
        self.assertEqual(data["mode"], "SENSIBULL PRACTICE")
        self.assertEqual(data["data_status"], "CSV SNAPSHOT")
        self.assertEqual(data["market"]["spot"], 22620.85)
        self.assertEqual(data["market"]["vix_source"], "Config.xlsx (manual)")
        self.assertEqual(data["market"]["futures"], 22676.34)
        self.assertEqual(len(data["bull_put"]), 13)
        self.assertEqual(len(data["bear_call"]), 9)
        self.assertEqual(data["bull_put"][0]["SellStrike"], 21750)
        self.assertEqual(data["bear_call"][0]["SellStrike"], 23400)
        self.assertEqual(data["bull_put"][0]["Rank"], 1)
        self.assertEqual(data["bear_call"][0]["Rank"], 1)
        self.assertIn("SellIV", data["bull_put"][0])
        self.assertIn("LiquidityScore", data["bull_put"][0])
        self.assertEqual(len(data["option_chain"]), 81)
        self.assertEqual((self.input / "Config.xlsx").read_bytes(), original_config)
        for route in ("/", "/dashboard", "/journal", "/settings", "/api/dashboard", "/api/health", "/downloads/journal", "/downloads/results"):
            with self.client.get(route) as page:
                self.assertEqual(page.status_code, 200, route)
        html = self.client.get("/dashboard").get_data(as_text=True)
        self.assertIn("22,620.85", html)
        self.assertIn("CSV SNAPSHOT", html)
        self.assertIn("Rank #1 within strategy", html)
        self.assertIn("id=\"chainTable\"", html)
        self.assertNotIn("cdn.jsdelivr", html)

    def test_journal_refresh_displays_saved_changes_without_scanning_or_writing(self):
        self.assertEqual(self.scan().status_code, 200)
        journal = self.output / "journal_tracker.xlsx"
        workbook = load_workbook(journal)
        workbook["Journal"]["L2"] = 12345.67
        workbook.save(journal)
        workbook.close()
        journal_before = journal.read_bytes()
        result_before = (self.output / "result.json").read_bytes()
        response = self.client.get("/journal")
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn('id="refreshJournalBtn" type="submit">Refresh Journal</button>', html)
        self.assertIn('method="get" action="/journal"', html)
        self.assertIn("12,345.67", html)
        self.assertEqual(html.count('data-trade-status="closed"'), 0)
        self.assertEqual(html.count('data-trade-status="open"'), 4)
        self.assertEqual(journal.read_bytes(), journal_before)
        self.assertEqual((self.output / "result.json").read_bytes(), result_before)
        self.assertIn("no-store", response.headers["Cache-Control"])

    def test_journal_kpis_follow_status_and_preserve_workbook(self):
        journal = self.output / "journal_tracker.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Journal"
        sheet.append(["Rank", "Loss", "Standalone Funds", "Profit", "Standalone Margin"])
        sheet.append([1, 0, 999999, 1250.25, 45000.50])
        sheet.append([2, -200.50, 888888, None, 30000])
        sheet.append([3, None, None, 49.75, None])
        workbook.save(journal)
        workbook.close()
        original = journal.read_bytes()
        runner = Mock(side_effect=AssertionError("KPI changes must not scan"))
        client = create_app(self.manager, analysis_runner=runner).test_client()
        html = client.get("/journal").get_data(as_text=True)
        self.assertNotIn("Continue editing the original workbook", html)
        self.assertEqual(html.count('data-journal-kpi="'), 5)
        self.assertIn('data-journal-kpi="capital">₹300,000.00', html)
        self.assertIn('data-journal-kpi="total_profit">₹1,300.00', html)
        self.assertIn('data-journal-kpi="total_loss">₹200.50', html)
        self.assertIn('data-journal-kpi="used_capital">₹75,000.50', html)
        self.assertIn('data-journal-kpi="available_capital">₹226,299.50', html)
        ids = re.findall(r'data-entry-id="([^"]+)"', html)
        response = client.post("/api/journal/status", json={"entry_id": ids[0], "status": "closed"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["kpis"], {
            "capital": 300000, "total_profit": 1300, "total_loss": 200.50,
            "used_capital": 30000, "available_capital": 271300})
        restarted = create_app(self.manager, analysis_runner=runner).test_client()
        self.assertIn('data-journal-kpi="used_capital">₹30,000.00', restarted.get("/journal").get_data(as_text=True))
        response = client.post("/api/journal/status", json={"entry_id": ids[0], "status": "open"})
        self.assertEqual(response.get_json()["kpis"]["used_capital"], 75000.50)
        with patch("journal_state.atomic_json", side_effect=PermissionError("Locked")):
            failed = client.post("/api/journal/status", json={"entry_id": ids[0], "status": "closed"})
        self.assertEqual(failed.status_code, 409)
        self.assertNotIn("kpis", failed.get_json())
        self.assertIn('data-journal-kpi="used_capital">₹75,000.50', client.get("/journal").get_data(as_text=True))
        self.assertEqual(journal.read_bytes(), original)
        runner.assert_not_called()
        self.assertFalse((self.output / "result.json").exists())

    def test_journal_kpis_refresh_saved_manual_amounts(self):
        self.scan()
        journal = self.output / "journal_tracker.xlsx"
        workbook = load_workbook(journal)
        workbook["Journal"]["K2"] = 60000
        workbook["Journal"]["L2"] = 3000
        workbook["Journal"]["M3"] = 500
        workbook.save(journal)
        workbook.close()
        original = journal.read_bytes()
        html = self.client.get("/journal").get_data(as_text=True)
        self.assertIn('data-journal-kpi="available_capital">₹243,000.00', html)
        self.assertIn('data-journal-kpi="total_loss">₹500.00', html)
        response = self.set_journal_status(self.journal_ids()[0], "closed")
        self.assertEqual(response.get_json()["kpis"]["available_capital"], 303000)
        self.assertEqual(journal.read_bytes(), original)

    def test_journal_kpis_blank_empty_invalid_and_unknown_values(self):
        headers = [" Profit ", "LOSS", "Standalone Margin"]
        empty = journal_kpis([], [], [])
        self.assertEqual(empty, {"capital": 300000, "total_profit": 0, "total_loss": 0,
                                "used_capital": 0, "available_capital": 300000})
        self.assertEqual(journal_kpis(headers, [[None, " ", None]], ["open"]), empty)
        self.assertEqual(journal_kpis(headers, [[0.1, 0, 0], [0.2, 0, 0]], ["open", "open"])["total_profit"], 0.3)
        self.assertEqual(journal_kpis(headers, [[0, 0, 400000]], ["open"])["available_capital"], -100000)
        for value in ("=1+2", "bad", "100", True, float("nan"), float("inf"), -1):
            with self.subTest(value=value):
                kpis = journal_kpis(headers, [[value, 0, 0]], ["open"])
                self.assertIsNone(kpis["total_profit"])
                self.assertIsNone(kpis["available_capital"])
                self.assertEqual(kpis["used_capital"], 0)
                kpis = journal_kpis(headers, [[0, 0, value]], ["open"])
                self.assertIsNone(kpis["used_capital"])
                self.assertIsNone(kpis["available_capital"])
                self.assertEqual(journal_kpis(headers, [[0, 0, value]], ["closed"])["used_capital"], 0)
        for key in ("used_capital", "available_capital"):
            self.assertIsNone(journal_kpis(headers, [[0, 0, 100]], ["unknown"])[key])
        missing = journal_kpis(["Rank"], [[1]], ["open"])
        self.assertTrue(all(missing[key] is None for key in missing if key != "capital"))
        unavailable = journal_kpis([], [], [], workbook_available=False)
        self.assertTrue(all(unavailable[key] is None for key in unavailable if key != "capital"))

    def test_journal_unreadable_workbook_kpis_do_not_claim_zero(self):
        (self.output / "journal_tracker.xlsx").write_bytes(b"invalid workbook")
        html = self.client.get("/journal").get_data(as_text=True)
        self.assertIn("Unable to read the journal workbook", html)
        for key in ("total_profit", "total_loss", "available_capital", "used_capital"):
            self.assertIn(f'data-journal-kpi="{key}">—', html)
        self.assertIn('data-journal-kpi="capital">₹300,000.00', html)

    def test_journal_default_status_ignores_profit_and_loss(self):
        journal = self.output / "journal_tracker.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Journal"
        sheet.append(["Rank", "Profit", "Loss"])
        for row in ([1, None, None], [2, 975, 0], [3, 0, 200], [4, 0, 0],
                    [5, 0, None], [6, " ", " "]):
            sheet.append(row)
        workbook.save(journal)
        workbook.close()
        original = journal.read_bytes()
        html = self.client.get("/journal").get_data(as_text=True)
        self.assertIn("<th>Trade status</th>", html)
        self.assertEqual(html.count('data-trade-status="open"'), 6)
        self.assertEqual(html.count('role="switch"'), 6)
        self.assertIn("6 Open", html)
        self.assertIn("0 Closed", html)
        self.assertEqual(self.set_journal_status(self.journal_ids()[0], "closed").status_code, 200)
        html = self.client.get("/journal").get_data(as_text=True)
        self.assertEqual(html.count('data-trade-status="closed"'), 1)
        self.assertEqual(journal.read_bytes(), original)

    def test_journal_toggle_persists_refresh_restart_without_editing_excel(self):
        self.assertEqual(self.scan().status_code, 200)
        ids = self.journal_ids()
        journal = self.output / "journal_tracker.xlsx"
        original = journal.read_bytes()
        result_before = (self.output / "result.json").read_bytes()
        config_before = (self.input / "Config.xlsx").read_bytes()
        response = self.set_journal_status(ids[0], "closed")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["counts"], {"open": 3, "closed": 1})
        self.assertEqual(self.set_journal_status(ids[1], "closed").status_code, 200)
        reopened = create_app(self.manager, analysis_runner=Mock(side_effect=AssertionError("No scan")))
        html = reopened.test_client().get("/journal").get_data(as_text=True)
        self.assertEqual(html.count('data-trade-status="closed"'), 2)
        self.assertEqual(html.count('data-trade-status="open"'), 2)
        response = self.set_journal_status(ids[0], "open")
        self.assertEqual(response.get_json()["counts"], {"open": 3, "closed": 1})
        saved = json.loads((self.output / "journal_status.json").read_text(encoding="utf-8"))
        self.assertEqual(saved["statuses"], {ids[0]: "open", ids[1]: "closed"})
        self.assertEqual(journal.read_bytes(), original)
        self.assertEqual((self.output / "result.json").read_bytes(), result_before)
        self.assertEqual((self.input / "Config.xlsx").read_bytes(), config_before)

    def test_journal_manual_status_is_unaffected_by_saved_pnl_changes(self):
        self.assertEqual(self.scan().status_code, 200)
        ids = self.journal_ids()
        self.set_journal_status(ids[0], "closed")
        journal = self.output / "journal_tracker.xlsx"
        workbook = load_workbook(journal)
        workbook["Journal"]["L2"] = 1500
        workbook["Journal"]["M3"] = 1000
        workbook.save(journal)
        workbook.close()
        self.assertEqual(self.journal_ids(), ids)
        html = self.client.get("/journal").get_data(as_text=True)
        self.assertEqual(html.count('data-trade-status="closed"'), 1)
        self.assertEqual(html.count('data-trade-status="open"'), 3)
        workbook = load_workbook(journal)
        workbook["Journal"]["L2"] = None
        workbook.save(journal)
        workbook.close()
        html = self.client.get("/journal").get_data(as_text=True)
        self.assertIn(f'data-entry-id="{ids[0]}" data-trade-status="closed"', html)

    def test_manual_status_works_without_pnl_columns(self):
        journal = self.output / "journal_tracker.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Journal"
        sheet.append(["Rank", "Notes"])
        sheet.append([1, "Practice trade"])
        workbook.save(journal)
        workbook.close()
        entry_id = self.journal_ids()[0]
        self.assertEqual(self.set_journal_status(entry_id, "closed").status_code, 200)
        html = self.client.get("/journal").get_data(as_text=True)
        self.assertEqual(html.count('data-trade-status="closed"'), 1)

    def test_repeat_scan_retains_manual_status_and_new_entries_default_open(self):
        self.scan()
        ids = self.journal_ids()
        self.set_journal_status(ids[0], "closed")
        state_before = (self.output / "journal_status.json").read_bytes()
        self.assertEqual(self.scan().get_json()["journal_rows_added"], 0)
        self.assertEqual(self.journal_ids(), ids)
        self.assertEqual((self.output / "journal_status.json").read_bytes(), state_before)
        journal = self.output / "journal_tracker.xlsx"
        workbook = load_workbook(journal)
        sheet = workbook["Journal"]
        row = [cell.value for cell in sheet[2]]
        row[6] += 1  # A separate observation, only in the temporary test workbook.
        sheet.append(row)
        workbook.save(journal)
        workbook.close()
        html = self.client.get("/journal").get_data(as_text=True)
        self.assertEqual(html.count('data-trade-status="closed"'), 1)
        self.assertEqual(html.count('data-trade-status="open"'), 4)

    def test_reordering_journal_rows_keeps_selection_with_trade(self):
        self.scan()
        ids = self.journal_ids()
        self.set_journal_status(ids[0], "closed")
        journal = self.output / "journal_tracker.xlsx"
        workbook = load_workbook(journal)
        sheet = workbook["Journal"]
        row = [cell.value for cell in sheet[2]]
        sheet.delete_rows(2)
        row[0] = 7  # A rank edit must not change the trade identity either.
        sheet.append(row)
        workbook.save(journal)
        workbook.close()
        self.assertEqual(self.journal_ids()[-1], ids[0])
        html = self.client.get("/journal").get_data(as_text=True)
        self.assertIn(f'data-entry-id="{ids[0]}" data-trade-status="closed"', html)

    def test_journal_status_rejects_invalid_and_stale_requests(self):
        self.scan()
        ids = self.journal_ids()
        for body in (None, [], {}, {"entry_id": ids[0], "status": True},
                     {"entry_id": ids[0], "status": "pending"}, {"entry_id": [], "status": "closed"}):
            response = self.client.post("/api/journal/status", json=body)
            self.assertEqual(response.status_code, 400)
        self.assertEqual(self.set_journal_status("not-a-journal-entry", "closed").status_code, 404)
        self.assertFalse((self.output / "journal_status.json").exists())

    def test_malformed_saved_status_disables_controls_without_overwriting(self):
        self.scan()
        ids = self.journal_ids()
        path = self.output / "journal_status.json"
        path.write_text("{broken", encoding="utf-8")
        original = path.read_bytes()
        html = self.client.get("/journal").get_data(as_text=True)
        self.assertEqual(html.count('data-trade-status="unknown"'), 4)
        self.assertEqual(html.count('title="On: Open · Off: Closed" disabled'), 4)
        self.assertEqual(self.set_journal_status(ids[0], "closed").status_code, 409)
        self.assertEqual(path.read_bytes(), original)

    def test_status_save_permission_failure_preserves_previous_selection(self):
        self.scan()
        ids = self.journal_ids()
        self.set_journal_status(ids[0], "closed")
        path = self.output / "journal_status.json"
        original = path.read_bytes()
        with patch("journal_state.atomic_json", side_effect=PermissionError("Locked")):
            response = self.set_journal_status(ids[0], "open")
        self.assertEqual(response.status_code, 409)
        self.assertIn("Unable to save", response.get_json()["error"])
        self.assertEqual(path.read_bytes(), original)

    def test_repeat_scan_is_idempotent_and_preserves_manual_entries(self):
        self.assertEqual(self.scan().status_code, 200)
        journal = self.output / "journal_tracker.xlsx"
        workbook = load_workbook(journal)
        sheet = workbook["Journal"]
        sheet["J2"] = 51000
        sheet["K2"] = 48000
        sheet["L2"] = 975
        sheet["M2"] = 0
        workbook.save(journal)
        workbook.close()
        before = journal.read_bytes()
        response = self.scan()
        self.assertEqual(response.get_json()["journal_rows_added"], 0)
        self.assertEqual(journal.read_bytes(), before)
        workbook = load_workbook(journal, read_only=True)
        self.assertEqual(workbook["Journal"].max_row, 5)
        workbook.close()

    def test_append_preserves_existing_layout_formulas_styles_extra_sheets(self):
        journal = self.output / "journal_tracker.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Journal"
        # The user's current layout intentionally has no Available Fund column.
        headers = TradeTracker.JOURNAL_COLUMNS[:-1] + ["Notes"]
        sheet.append(headers)
        sheet.append([1, "2026-08-25", "Bull Put", "PE", 24000, 23700, 54.55, 24.65, 83,
                      51107, 47561, "=10+20", 0, "Keep my notes"])
        sheet["L2"].font = Font(bold=True, color="00AA00")
        sheet.freeze_panes = "D2"
        workbook.create_sheet("My calculations")["A1"] = "=SUM(Journal!J:J)"
        workbook.save(journal)
        workbook.close()
        original = journal.read_bytes()
        response = self.scan()
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(journal)
        sheet = workbook["Journal"]
        self.assertEqual([cell.value for cell in sheet[1]], headers)
        self.assertEqual(sheet.max_row, 6)
        self.assertEqual(sheet["J2"].value, 51107)
        self.assertEqual(sheet["L2"].value, "=10+20")
        self.assertTrue(sheet["L2"].font.bold)
        self.assertEqual(sheet["N2"].value, "Keep my notes")
        self.assertEqual(sheet.freeze_panes, "D2")
        self.assertEqual(workbook["My calculations"]["A1"].value, "=SUM(Journal!J:J)")
        appended = list(sheet.iter_rows(min_row=3, values_only=True))
        self.assertEqual([row[0] for row in appended], [1, 2, 1, 2])
        self.assertTrue(all(all(value is None for value in row[9:]) for row in appended))
        workbook.close()
        backups = list((self.output / "backups").glob("*.xlsx"))
        self.assertEqual(len(backups), 1)
        self.assertEqual(backups[0].read_bytes(), original)

    def test_no_trade_creates_empty_results_without_appending_journal(self):
        self.scan()
        journal = self.output / "journal_tracker.xlsx"
        before = journal.read_bytes()
        self.manager.update_config_values({"MinPOP": 100})
        response = self.scan()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["journal_rows_added"], 0)
        self.assertEqual(self.load_result()["bull_put"], [])
        self.assertEqual(self.load_result()["bear_call"], [])
        self.assertEqual(journal.read_bytes(), before)
        self.assertIn("NO TRADE", self.client.get("/dashboard").get_data(as_text=True))

    def test_invalid_requests_never_call_analysis(self):
        runner = Mock()
        client = create_app(self.manager, runner).test_client()
        bodies = [None, [], {}, {"expiry": 123}, {"expiry": "bad"},
                  {"expiry": "2000-01-01"}, {"expiry": self.expiry, "market_source": "unknown"}]
        for body in bodies:
            with self.subTest(body=body):
                response = client.post("/run-analysis", data=json.dumps(body), content_type="application/json")
                self.assertEqual(response.status_code, 400)
                self.assertFalse(response.get_json()["success"])
        runner.assert_not_called()

    def test_expired_upload_not_selected(self):
        expired = self.input / "NIFTY_2000-01-01_option_chain_2000-01-01-10-00-00.csv"
        shutil.copy2(self.csv, expired)
        self.assertEqual(self.manager.find_option_chain().name, self.csv.name)
        self.assertEqual(self.scan().status_code, 200)
        result = self.load_result()
        result["expiry_date"] = "2000-01-01"
        (self.output / "result.json").write_text(json.dumps(result), encoding="utf-8")
        self.assertEqual(self.client.get("/api/dashboard").get_json()["status"], "EXPIRED SNAPSHOT")

    def test_latest_download_for_same_expiry_is_selected(self):
        older = self.input / f"NIFTY_{self.expiry}_option_chain_2020-01-01-09-00-00.csv"
        shutil.copy2(self.csv, older)
        self.assertEqual(self.manager.find_option_chain(date.fromisoformat(self.expiry)), self.csv)

    def test_duplicate_strikes_rejected_and_old_result_retained(self):
        self.scan()
        result = (self.output / "result.json").read_bytes()
        frame = pd.read_csv(self.csv)
        pd.concat([frame, frame.head(1)]).to_csv(self.csv, index=False)
        response = self.scan()
        self.assertEqual(response.status_code, 400)
        self.assertIn("duplicate strikes", response.get_json()["error"])
        self.assertEqual((self.output / "result.json").read_bytes(), result)

    def test_all_missing_numeric_column_rejected_as_candidates_not_exception(self):
        frame = pd.read_csv(self.csv)
        frame["Call Volume"] = "--"
        frame.to_csv(self.csv, index=False)
        response = self.scan()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.load_result()["bear_call"], [])

    def test_missing_unused_pop_columns_does_not_crash(self):
        frame = pd.read_csv(self.csv).drop(columns=["Call POP %", "Put POP %"])
        frame.to_csv(self.csv, index=False)
        self.assertEqual(self.scan().status_code, 200)

    def test_csv_inconsistent_recorded_spot_rejected(self):
        frame = pd.read_csv(self.csv)
        column = "Call Intrinsic Value(Spot)"
        index = frame.index[pd.to_numeric(frame[column], errors="coerce").gt(0)][0]
        frame.loc[index, column] = float(frame.loc[index, column]) + 100
        frame.to_csv(self.csv, index=False)
        response = self.scan()
        self.assertEqual(response.status_code, 400)
        self.assertIn("inconsistent spot", response.get_json()["error"])

    def test_malformed_result_is_visible_not_silent_empty_success(self):
        path = self.output / "result.json"
        for contents in ("{bad", "[]", '{"bull_put":[{}]}', '{"bull_put":[{"Rank":NaN}]}'):
            path.write_text(contents, encoding="utf-8")
            response = self.client.get("/api/dashboard")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.get_json()["status"], "DATA ERROR")
            self.assertTrue(response.get_json()["error"])
            self.assertEqual(self.client.get("/dashboard").status_code, 200)

    def test_excel_locked_error_and_previous_results_retained(self):
        self.scan()
        result = (self.output / "result.json").read_bytes()
        with patch("trade_tracker.atomic_workbook", side_effect=PermissionError("locked")):
            # Force one additional observation so the journal needs a write.
            self.manager.update_config_values({"WingWidth": 250})
            response = self.scan()
        self.assertEqual(response.status_code, 409)
        self.assertIn("Close", response.get_json()["error"])
        self.assertEqual((self.output / "result.json").read_bytes(), result)

    def test_optional_broker_failure_keeps_prior_result(self):
        self.scan()
        before = (self.output / "result.json").read_bytes()
        with patch("data_provider.fetch_spot_and_vix", side_effect=MarketDataError("SmartAPI authentication failed.")):
            response = self.scan("live")
        self.assertEqual(response.status_code, 502)
        self.assertEqual((self.output / "result.json").read_bytes(), before)

    def test_concurrent_scan_returns_conflict(self):
        client = create_app(self.manager, Mock(side_effect=AnalysisBusy("Already running"))).test_client()
        response = client.post("/run-analysis", json={"expiry": self.expiry})
        self.assertEqual(response.status_code, 409)

    def test_rulebook_reloaded_on_every_scan(self):
        self.scan()
        self.manager.update_config_values({"MinPOP": 90})
        self.scan()
        data = self.load_result()
        self.assertEqual(data["rules"]["MinPOP"], 90)
        self.assertTrue(all(row["EstimatedPOP"] >= 90 for row in data["bull_put"] + data["bear_call"]))

    def test_downloads_do_not_accept_arbitrary_paths(self):
        response = self.client.get("/downloads/../../Input/Config.xlsx")
        self.assertEqual(response.status_code, 404)

    def test_configuration_invalid_delta_rejected(self):
        self.manager.update_config_values({"MinDelta": .8, "MaxDelta": .4})
        self.assertEqual(self.scan().status_code, 400)

    def test_csv_scan_does_not_update_configuration_workbook(self):
        before = hashlib.sha256((self.input / "Config.xlsx").read_bytes()).hexdigest()
        run_analysis(self.expiry, config_manager=self.manager)
        self.assertEqual(hashlib.sha256((self.input / "Config.xlsx").read_bytes()).hexdigest(), before)

    def test_new_journal_available_fund_formula_and_top_two_only(self):
        self.scan()
        workbook = load_workbook(self.output / "journal_tracker.xlsx")
        sheet = workbook["Journal"]
        self.assertEqual(sheet.max_row, 5)
        self.assertEqual(sheet["N2"].value, "=300000.0-J2")
        self.assertEqual(sheet["N3"].value, "=N2-J3")
        workbook.close()


class CalculationTests(unittest.TestCase):
    def setUp(self):
        self.context = MarketContext(24000, 15, "Sideways", "Wide", date.today(), date.today()+timedelta(days=14),
                                     14, 65, 300, 300000, 70, .1, .4, 10000, 5000, 3, 2, 250000)
        self.engine = StrategyEngine(pd.DataFrame(), self.context)

    def spread(self, strategy="Bear Call"):
        return pd.DataFrame([{"Strategy": strategy, "SellStrike": 24200 if strategy == "Bear Call" else 23800,
                              "BuyStrike": 24500 if strategy == "Bear Call" else 23500,
                              "SellPremium": 64.7, "BuyPremium": 23.4,
                              "SellDelta": .23 if strategy == "Bear Call" else -.23,
                              "BuyDelta": .10 if strategy == "Bear Call" else -.10,
                              "SellTheta": -4.2, "BuyTheta": -2.22,
                              "SellVega": 8, "BuyVega": 4,
                              "SellOI": 30000, "BuyOI": 50000}])

    def test_manual_bear_call_payoff_and_greeks(self):
        row = self.engine.calculate_quality_metrics(self.engine.calculate_metrics(self.spread())).iloc[0]
        expected = {"NetCredit":41.3, "SpreadWidth":300, "MaxProfit":2684.5, "MaxLoss":16815.5,
                    "BreakEven":24241.3, "ReturnOnRisk":15.96, "EstimatedPOP":77,
                    "NetDelta":-.13, "NetTheta":1.98, "NetVega":-4,
                    "LiquidityScore":40000, "SafetyMargin":241.3}
        for key, value in expected.items():
            self.assertAlmostEqual(row[key], value, places=2, msg=key)
        self.assertAlmostEqual(row.CreditEfficiency, 41.3/300*100)

    def test_manual_bull_put_breakeven_and_greeks(self):
        row = self.engine.calculate_metrics(self.spread("Bull Put")).iloc[0]
        self.assertAlmostEqual(row.BreakEven, 23758.7)
        self.assertAlmostEqual(row.NetDelta, .13)
        self.assertAlmostEqual(row.NetTheta, 1.98)

    def test_degenerate_credit_risk_never_eligible(self):
        for credit in (0, -1, 300, 301, np.nan):
            raw = self.spread()
            raw["SellPremium"] = raw["BuyPremium"] + credit
            calculated = self.engine.calculate_metrics(raw)
            self.assertTrue(self.engine.filter_strategies(calculated).empty)

    def test_ranking_preserves_original_priority_order(self):
        frame = pd.DataFrame([
            {"EstimatedPOP":80,"CreditEfficiency":10,"ReturnOnRisk":30,"SafetyMargin":200,"LiquidityScore":1000},
            {"EstimatedPOP":90,"CreditEfficiency":5,"ReturnOnRisk":10,"SafetyMargin":100,"LiquidityScore":10},
            {"EstimatedPOP":80,"CreditEfficiency":20,"ReturnOnRisk":5,"SafetyMargin":50,"LiquidityScore":50},
        ])
        ranked = self.engine.rank_strategies(frame)
        self.assertEqual(ranked.EstimatedPOP.tolist(), [90,80,80])
        self.assertEqual(ranked.CreditEfficiency.tolist(), [5,20,10])
        self.assertEqual(ranked.Rank.tolist(), [1,2,3])

    def test_spread_points_rule_uses_points_not_percentage(self):
        frame = pd.DataFrame({"CallBid":[9],"CallAsk":[10],"PutBid":[9],"PutAsk":[10],
                              "CallOI":[20000],"PutOI":[20000],"CallVolume":[6000],"PutVolume":[6000]})
        scorer = OptionScorer(frame, self.context)
        scorer.calculate_spread()
        scorer.calculate_liquidity()
        self.assertEqual(frame.CallSpreadPct.iloc[0], 10)
        self.assertTrue(frame.CallLiquid.iloc[0])  # 1 point <= 2 points, despite 10% spread.

    def test_crossed_quote_and_missing_hedge_data_rejected(self):
        row = {"IV":15, "CallLTP":20, "CallBid":19, "CallAsk":21, "CallDelta":.1,
               "CallTheta":-2, "CallVega":4, "CallOI":10000, "CallVolume":5000}
        self.assertTrue(self.engine._valid_leg(pd.Series(row), "Call"))
        for changes in ({"CallBid":22},{"CallLTP":0},{"CallOI":np.nan},{"CallVolume":np.nan},{"CallDelta":-2}):
            self.assertFalse(self.engine._valid_leg(pd.Series({**row, **changes}), "Call"))


if __name__ == "__main__":
    unittest.main()
